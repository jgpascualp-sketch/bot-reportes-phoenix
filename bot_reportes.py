import os
import sys
import logging
import threading
from datetime import datetime
from http.server import HTTPServer, BaseHTTPRequestHandler
from telegram import Update, ReplyKeyboardMarkup, ReplyKeyboardRemove, InlineKeyboardMarkup, InlineKeyboardButton
from telegram.ext import (
    ApplicationBuilder,
    CommandHandler,
    MessageHandler,
    CallbackQueryHandler,
    ConversationHandler,
    ContextTypes,
    filters,
)
import docx
from docx.shared import Inches, Pt
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_ALIGN_VERTICAL
from docx.oxml import parse_xml
from docx.oxml.ns import nsdecls

logging.basicConfig(
    format="%(asctime)s - %(name)s - %(levelname)s - %(message)s",
    level=logging.INFO,
    stream=sys.stdout
)
logger = logging.getLogger(__name__)

# Servidor HTTP Render
class HealthHandler(BaseHTTPRequestHandler):
    def do_GET(self):
        self.send_response(200)
        self.send_header("Content-type", "text/plain")
        self.end_headers()
        self.wfile.write(b"OK - Bot Activo")

    def log_message(self, format, *args):
        pass

def iniciar_servidor_web():
    puerto = int(os.environ.get("PORT", 10000))
    servidor = HTTPServer(("0.0.0.0", puerto), HealthHandler)
    servidor.serve_forever()

ITEMS_CHECKLIST = [
    "Apariencia (Appearance check)",
    "Bateria de respaldo (Backup battery)",
    "Placa Hall o Tarjeta electronica",
    "Señal de sensor (Sensor signal)",
    "Pantalla táctil (Touch screen)",
    "Pieza hidráulica (Hydraulic parts)",
    "Parametros de Tratamiento",
    "Sim. de Tratamiento (Simulation treatme)",
    "Opciones (Options)",
    "Sensor de Cond. (Conductivity sensor)",
    "Bomba Ceramica (Ceramic pump)",
    "Bomba de Heparina (Syringe pump)",
    "Calibración de Conductividad (Cond. Calibration)",
    "Calibración de Temperatura (Temp. calibration)",
    "Calibración de Presión (Pressure calibration)",
    "Otros (Other)"
]

(
    CONSECUTIVO,
    HOSPITAL,
    CONTACTO,
    TELEFONO,
    DIRECCION,
    MODELO,
    SERIE,
    HOROMETRO,
    VERSION_SW,
    TIPO_SERVICIO,
    DETALLES,
    FALLA_TIPO,
    SOLUCION,
    CHECKLIST_MENU,
    REP_PARTE,
    REP_CANT,
    REP_OBS,
    SATISFACCION,
    INGENIERO,
    OPCION_FIRMA,
    SUBIR_FIRMA,
    FECHA,
    MONEDA,
) = range(23)

async def start(update: Update, context: ContextTypes.DEFAULT_TYPE):
    await update.message.reply_text("👋 Usa /reporte para generar un nuevo reporte técnico.")

async def iniciar_reporte(update: Update, context: ContextTypes.DEFAULT_TYPE):
    context.user_data.clear()
    context.user_data["checklist_seleccionados"] = []
    teclado = [["Dejar vacío"]]
    reply_markup = ReplyKeyboardMarkup(teclado, one_time_keyboard=True, resize_keyboard=True)
    await update.message.reply_text("📄 Ingrese Consecutivo (ej: 0000008) o pulse Dejar vacío:", reply_markup=reply_markup)
    return CONSECUTIVO

async def get_consecutivo(update: Update, context: ContextTypes.DEFAULT_TYPE):
    txt = update.message.text
    context.user_data["consecutivo"] = "" if txt == "Dejar vacío" else txt
    await update.message.reply_text("🏥 Ingrese Nombre de la Clínica / Hospital:", reply_markup=ReplyKeyboardRemove())
    return HOSPITAL

async def get_hospital(update: Update, context: ContextTypes.DEFAULT_TYPE):
    context.user_data["hospital"] = update.message.text
    await update.message.reply_text("👤 Ingrese Nombre del Contacto:")
    return CONTACTO

async def get_contacto(update: Update, context: ContextTypes.DEFAULT_TYPE):
    context.user_data["contacto"] = update.message.text
    await update.message.reply_text("📞 Ingrese Número de Teléfono:")
    return TELEFONO

async def get_telefono(update: Update, context: ContextTypes.DEFAULT_TYPE):
    context.user_data["telefono"] = update.message.text
    await update.message.reply_text("📍 Ingrese Dirección:")
    return DIRECCION

async def get_direccion(update: Update, context: ContextTypes.DEFAULT_TYPE):
    context.user_data["direccion"] = update.message.text
    teclado = [["DORA-6000"], ["Otro"]]
    reply_markup = ReplyKeyboardMarkup(teclado, one_time_keyboard=True, resize_keyboard=True)
    await update.message.reply_text("🤖 Seleccione o ingrese Modelo:", reply_markup=reply_markup)
    return MODELO

async def get_modelo(update: Update, context: ContextTypes.DEFAULT_TYPE):
    context.user_data["modelo"] = update.message.text
    await update.message.reply_text("🔢 Ingrese Serial No.:", reply_markup=ReplyKeyboardRemove())
    return SERIE

async def get_serie(update: Update, context: ContextTypes.DEFAULT_TYPE):
    context.user_data["serie"] = update.message.text
    teclado = [["Dejar vacío"]]
    reply_markup = ReplyKeyboardMarkup(teclado, one_time_keyboard=True, resize_keyboard=True)
    await update.message.reply_text("⏱️ Ingrese Horómetro o pulse Dejar vacío:", reply_markup=reply_markup)
    return HOROMETRO

async def get_horometro(update: Update, context: ContextTypes.DEFAULT_TYPE):
    txt = update.message.text
    context.user_data["horometro"] = "" if txt == "Dejar vacío" else txt
    teclado = [["020305"], ["Dejar vacío"]]
    reply_markup = ReplyKeyboardMarkup(teclado, one_time_keyboard=True, resize_keyboard=True)
    await update.message.reply_text("💻 Ingrese Versión de Software:", reply_markup=reply_markup)
    return VERSION_SW

async def get_version_sw(update: Update, context: ContextTypes.DEFAULT_TYPE):
    txt = update.message.text
    context.user_data["version_sw"] = "" if txt == "Dejar vacío" else txt
    teclado = [
        ["Mantenimiento Correctivo (Repair)"],
        ["Mantenimiento Preventivo (PM)"],
        ["Instalación (Installation)"],
        ["Diagnostico (Diagnostic / Inspection)"],
        ["Entrenamiento (Operation training)"],
        ["Seguimiento Tratamiento (Follow-up Trearment)"],
        ["Desinstalación (Uninstallation)"],
        ["Otro (Other)"]
    ]
    reply_markup = ReplyKeyboardMarkup(teclado, one_time_keyboard=True, resize_keyboard=True)
    await update.message.reply_text("🛠️ Seleccione Tipo de Servicio:", reply_markup=reply_markup)
    return TIPO_SERVICIO

async def get_tipo_servicio(update: Update, context: ContextTypes.DEFAULT_TYPE):
    context.user_data["tipo_servicio"] = update.message.text
    await update.message.reply_text("📝 Ingrese Detalles / Feedback Details:", reply_markup=ReplyKeyboardRemove())
    return DETALLES

async def get_detalles(update: Update, context: ContextTypes.DEFAULT_TYPE):
    context.user_data["detalles"] = update.message.text
    teclado = [
        ["Fallo Hidaulico (Hydraulic fault)"],
        ["Fallo en el Circuito (Circuit fault)"],
        ["Fallo en parte de sangre (Bloodparts fault)"],
        ["Fallo en Software (Software fault)"],
        ["Fallo Mecanico (Mechanical fault)"],
        ["Fallo de montaje de pieza (Assemble fault)"],
        ["Fallo de desgaste rápido de pieza (Quick-wear part)"],
        ["Otros Fallos (Others fault)"],
        ["Ninguno / Normal"]
    ]
    reply_markup = ReplyKeyboardMarkup(teclado, one_time_keyboard=True, resize_keyboard=True)
    await update.message.reply_text("⚠️ Clasificación de Fallas:", reply_markup=reply_markup)
    return FALLA_TIPO

async def get_falla_tipo(update: Update, context: ContextTypes.DEFAULT_TYPE):
    context.user_data["falla_tipo"] = update.message.text
    teclado = [["Dejar vacío"]]
    reply_markup = ReplyKeyboardMarkup(teclado, one_time_keyboard=True, resize_keyboard=True)
    await update.message.reply_text("💡 Ingrese Motivo del Fallo y Solución:", reply_markup=reply_markup)
    return SOLUCION

def armar_teclado_checklist(seleccionados):
    botones = []
    fila = []
    for item in ITEMS_CHECKLIST:
        nombre_corto = item.split("(")[0].strip()
        marca = "✔ " if item in seleccionados else "⬜ "
        fila.append(InlineKeyboardButton(f"{marca}{nombre_corto}", callback_data=f"chk_{item}"))
        if len(fila) == 2:
            botones.append(fila)
            fila = []
    if fila:
        botones.append(fila)
    botones.append([InlineKeyboardButton("✨ SELECCIONAR TODO", callback_data="chk_ALL")])
    botones.append([InlineKeyboardButton("✅ LISTO / CONTINUAR", callback_data="chk_DONE")])
    return InlineKeyboardMarkup(botones)

async def get_solucion(update: Update, context: ContextTypes.DEFAULT_TYPE):
    txt = update.message.text
    context.user_data["solucion"] = "" if txt == "Dejar vacío" else txt
    
    markup = armar_teclado_checklist(context.user_data["checklist_seleccionados"])
    await update.message.reply_text(
        "📋 Lista de verificación de pruebas:\nSeleccione los ítems a marcar con [ ✔ ] y presione 'LISTO / CONTINUAR':",
        reply_markup=markup
    )
    return CHECKLIST_MENU

async def checklist_callback(update: Update, context: ContextTypes.DEFAULT_TYPE):
    query = update.callback_query
    await query.answer()
    data = query.data
    sel = context.user_data.get("checklist_seleccionados", [])
    
    if data == "chk_DONE":
        await query.message.reply_text(f"✅ Se seleccionaron {len(sel)} ítems.")
        teclado = [["Omitir / Todo Vacío"], ["Dejar vacío"]]
        reply_markup = ReplyKeyboardMarkup(teclado, one_time_keyboard=True, resize_keyboard=True)
        await query.message.reply_text("⚙️ Registro de componentes:\nIngrese Nombre de la parte (o pulse Omitir):", reply_markup=reply_markup)
        return REP_PARTE
        
    elif data == "chk_ALL":
        context.user_data["checklist_seleccionados"] = list(ITEMS_CHECKLIST)
    else:
        item = data.replace("chk_", "")
        if item in sel:
            sel.remove(item)
        else:
            sel.append(item)
        context.user_data["checklist_seleccionados"] = sel

    await query.edit_message_reply_markup(reply_markup=armar_teclado_checklist(context.user_data["checklist_seleccionados"]))
    return CHECKLIST_MENU

async def get_rep_parte(update: Update, context: ContextTypes.DEFAULT_TYPE):
    txt = update.message.text
    if txt == "Omitir / Todo Vacío":
        context.user_data["rep_parte"] = ""
        context.user_data["rep_cant"] = ""
        context.user_data["rep_obs"] = ""
        teclado = [
            ["Satisfecho (Satisfied)"],
            ["Relativamente satisfecho"],
            ["Normal (Normal)"],
            ["Insatisfecho (Dissatisfied)"],
            ["Muy Insatisfecho (Very Dissatisfied)"]
        ]
        reply_markup = ReplyKeyboardMarkup(teclado, one_time_keyboard=True, resize_keyboard=True)
        await update.message.reply_text("⭐ Encuesta de satisfacción / Opinión de usuario:", reply_markup=reply_markup)
        return SATISFACCION
        
    context.user_data["rep_parte"] = "" if txt == "Dejar vacío" else txt
    teclado = [["Dejar vacío"]]
    reply_markup = ReplyKeyboardMarkup(teclado, one_time_keyboard=True, resize_keyboard=True)
    await update.message.reply_text("⚙️ Ingrese Cantidad (Quantity):", reply_markup=reply_markup)
    return REP_CANT

async def get_rep_cant(update: Update, context: ContextTypes.DEFAULT_TYPE):
    txt = update.message.text
    context.user_data["rep_cant"] = "" if txt == "Dejar vacío" else txt
    teclado = [["Dejar vacío"]]
    reply_markup = ReplyKeyboardMarkup(teclado, one_time_keyboard=True, resize_keyboard=True)
    await update.message.reply_text("⚙️ Ingrese Observación (Remark):", reply_markup=reply_markup)
    return REP_OBS

async def get_rep_obs(update: Update, context: ContextTypes.DEFAULT_TYPE):
    txt = update.message.text
    context.user_data["rep_obs"] = "" if txt == "Dejar vacío" else txt
    teclado = [
        ["Satisfecho (Satisfied)"],
        ["Relativamente satisfecho"],
        ["Normal (Normal)"],
        ["Insatisfecho (Dissatisfied)"],
        ["Muy Insatisfecho (Very Dissatisfied)"]
    ]
    reply_markup = ReplyKeyboardMarkup(teclado, one_time_keyboard=True, resize_keyboard=True)
    await update.message.reply_text("⭐ Encuesta de satisfacción / Opinión de usuario:", reply_markup=reply_markup)
    return SATISFACCION

async def get_satisfaccion(update: Update, context: ContextTypes.DEFAULT_TYPE):
    context.user_data["satisfaccion"] = update.message.text
    teclado = [["Jesus Guillermo Pascual chalan"], ["Ingresar otro nombre"]]
    reply_markup = ReplyKeyboardMarkup(teclado, one_time_keyboard=True, resize_keyboard=True)
    await update.message.reply_text("👷 Ingeniero a cargo:", reply_markup=reply_markup)
    return INGENIERO

async def get_ingeniero(update: Update, context: ContextTypes.DEFAULT_TYPE):
    txt = update.message.text
    if txt == "Ingresar otro nombre":
        await update.message.reply_text("Escriba el nombre del ingeniero:", reply_markup=ReplyKeyboardRemove())
        return INGENIERO
    context.user_data["ingeniero"] = txt
    
    teclado = [["🖊️ Firma Automática (Jesús Pascual)"], ["📷 Subir Foto de Firma"]]
    reply_markup = ReplyKeyboardMarkup(teclado, one_time_keyboard=True, resize_keyboard=True)
    await update.message.reply_text("🖋️ ¿Cómo desea estampar la firma?", reply_markup=reply_markup)
    return OPCION_FIRMA

async def get_opcion_firma(update: Update, context: ContextTypes.DEFAULT_TYPE):
    txt = update.message.text
    if "Subir Foto" in txt:
        await update.message.reply_text("📸 Envíe la FOTO de la firma como imagen:", reply_markup=ReplyKeyboardRemove())
        return SUBIR_FIRMA
    
    context.user_data["firma_custom"] = None
    fecha_hoy = datetime.now().strftime("%d/%m/%Y")
    teclado = [[fecha_hoy], ["Ingresar otra fecha"]]
    reply_markup = ReplyKeyboardMarkup(teclado, one_time_keyboard=True, resize_keyboard=True)
    await update.message.reply_text("📅 Fecha del reporte (seleccione o edite):", reply_markup=reply_markup)
    return FECHA

async def get_foto_firma(update: Update, context: ContextTypes.DEFAULT_TYPE):
    foto = await update.message.photo[-1].get_file()
    ruta_temp = "firma_subida.png"
    await foto.download_to_drive(ruta_temp)
    context.user_data["firma_custom"] = ruta_temp
    
    fecha_hoy = datetime.now().strftime("%d/%m/%Y")
    teclado = [[fecha_hoy], ["Ingresar otra fecha"]]
    reply_markup = ReplyKeyboardMarkup(teclado, one_time_keyboard=True, resize_keyboard=True)
    await update.message.reply_text("✅ Firma recibida. 📅 Confirme la fecha del reporte:", reply_markup=reply_markup)
    return FECHA

async def get_fecha(update: Update, context: ContextTypes.DEFAULT_TYPE):
    txt = update.message.text
    if txt == "Ingresar otra fecha":
        await update.message.reply_text("Escriba la fecha (ej: 24/09/2026):", reply_markup=ReplyKeyboardRemove())
        return FECHA
    context.user_data["fecha"] = txt
    
    teclado = [["Dejar vacío"], ["Soles (S/.)"], ["USD ($)"]]
    reply_markup = ReplyKeyboardMarkup(teclado, one_time_keyboard=True, resize_keyboard=True)
    await update.message.reply_text("💵 Moneda y cobro:", reply_markup=reply_markup)
    return MONEDA

def escribir_con_espacio(celda, texto, negrita=False):
    celda.text = ""
    p = celda.paragraphs[0]
    p.paragraph_format.space_before = Pt(1)
    p.paragraph_format.space_after = Pt(1)
    run = p.add_run(f"  {texto}")
    run.font.size = Pt(8)
    if negrita:
        run.bold = True

def eliminar_linea_vertical(celda_izq, celda_der):
    tcPr_i = celda_izq._tc.get_or_add_tcPr()
    tcPr_d = celda_der._tc.get_or_add_tcPr()
    border_right_nil = parse_xml(f'<w:tcBorders {nsdecls("w")}><w:right w:val="nil"/></w:tcBorders>')
    border_left_nil = parse_xml(f'<w:tcBorders {nsdecls("w")}><w:left w:val="nil"/></w:tcBorders>')
    tcPr_i.append(border_right_nil)
    tcPr_d.append(border_left_nil)

async def get_moneda(update: Update, context: ContextTypes.DEFAULT_TYPE):
    try:
        txt = update.message.text
        context.user_data["moneda"] = "" if txt == "Dejar vacío" else txt
        
        plantilla = "1-TECHNICAL SERVICE REPORT corregido.docx"
        if not os.path.exists(plantilla):
            plantilla = "1-TECHNICAL SERVICE REPORT.docx"

        if not os.path.exists(plantilla):
            await update.message.reply_text("⚠️ No se encontró la plantilla en el repositorio.")
            return ConversationHandler.END

        doc = docx.Document(plantilla)
        
        consecutivo = context.user_data.get("consecutivo", "")
        hosp = context.user_data.get("hospital", "")
        contacto = context.user_data.get("contacto", "")
        telefono = context.user_data.get("telefono", "")
        direccion = context.user_data.get("direccion", "")
        modelo = context.user_data.get("modelo", "")
        serie = context.user_data.get("serie", "")
        horometro = context.user_data.get("horometro", "")
        version_sw = context.user_data.get("version_sw", "")
        detalles = context.user_data.get("detalles", "")
        solucion = context.user_data.get("solucion", "")
        tipo_serv = context.user_data.get("tipo_servicio", "")
        falla_tipo = context.user_data.get("falla_tipo", "")
        rep_p = context.user_data.get("rep_parte", "")
        rep_c = context.user_data.get("rep_cant", "")
        rep_o = context.user_data.get("rep_obs", "")
        satisfaccion = context.user_data.get("satisfaccion", "")
        checklist_sel = context.user_data.get("checklist_seleccionados", [])
        ingeniero = context.user_data.get("ingeniero", "Jesus Guillermo Pascual chalan")
        fecha_reporte = context.user_data.get("fecha", datetime.now().strftime("%d/%m/%Y"))

        t = doc.tables[0]

        # Consecutivo
        if consecutivo:
            for r in t.rows[:2]:
                for c in r.cells:
                    if "consecutivo" in c.text.lower():
                        c.text = f"  Consecutivo (Consecutive)\n  {consecutivo}"
                        break

        # 1. Datos cliente y equipo
        for r in t.rows:
            txt_fila = [c.text.strip().lower() for c in r.cells]
            if any("hospital name" in x for x in txt_fila):
                escribir_con_espacio(r.cells[1], hosp)
            if any("contact" in x for x in txt_fila) and any("phone" in x for x in txt_fila):
                escribir_con_espacio(r.cells[1], contacto)
                escribir_con_espacio(r.cells[-1], telefono)
            if any("adress" in x or "dirección" in x for x in txt_fila):
                escribir_con_espacio(r.cells[1], direccion)
            if any("model" in x for x in txt_fila) and any("version" in x for x in txt_fila):
                escribir_con_espacio(r.cells[1], modelo)
                escribir_con_espacio(r.cells[-1], version_sw)
            if any("serial no" in x for x in txt_fila):
                escribir_con_espacio(r.cells[1], serie)
            if any("running" in x or "horometro" in x for x in txt_fila):
                escribir_con_espacio(r.cells[1], horometro)
                
            # Detalles
            if any("feedback details" in x or "detalles" in x for x in txt_fila):
                r.cells[0].vertical_alignment = WD_ALIGN_VERTICAL.CENTER
                r.cells[0].text = ""
                p_det = r.cells[0].paragraphs[0]
                p_det.alignment = WD_ALIGN_PARAGRAPH.CENTER
                p_det.paragraph_format.space_before = Pt(0)
                p_det.paragraph_format.space_after = Pt(0)
                p_det.add_run("Detalles\n(Feedback Details)").font.size = Pt(8.5)
                escribir_con_espacio(r.cells[-1], detalles)
                
            if any("motivo del fallo" in x or "fault reason" in x for x in txt_fila):
                escribir_con_espacio(r.cells[-1], solucion)

        # 2. Tipo de Servicio
        col1 = [
            ("Mantenimiento Correctivo (Repair)", "correctivo"),
            ("Diagnostico (Diagnostic / Inspection)", "diagnostico"),
            ("Instalación (Installation)", "instal"),
            ("Desinstalación (Uninstallation)", "desinstal")
        ]
        col2 = [
            ("Mantenimiento Preventivo (PM)", "preventivo"),
            ("Entrenamiento (Operat
